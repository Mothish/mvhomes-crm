import csv
import io
from datetime import datetime, timedelta

import openpyxl
from django.contrib import messages
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db.models import Case, Q, When
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from ..access import (
    ADMIN, LEADER, MANAGER, MGMT, SALES_ROLES, assignable_users, can_view_lead, get_role, role_required,
    visible_leads, visible_users,
)
from ..forms import ActivityForm, ImportForm, LeadForm, LostForm, ReassignForm, strip_dial_code
from ..locale_data import (
    BASE_CURRENCY, CURRENCIES, DEFAULT_COUNTRY, DEFAULT_DIAL, country_code, dial_for, geo_data, to_base,
)
from ..messaging import lead_previews
from ..models import Lead, Project, normalize_phone
from ..services import assign, audit, capture_lead, log_activity, set_stage

BOARD_CLOSED_DAYS = 30


def get_visible_lead(request, pk):
    lead = get_object_or_404(Lead.objects.select_related("owner", "project"), pk=pk)
    if not can_view_lead(request.user, lead):
        raise PermissionDenied
    return lead


def filter_leads(request, queryset):
    """Shared filters for list, board and export. Returns (queryset, active-filter dict)."""
    g = request.GET
    f = {k: g.get(k, "").strip() for k in ("q", "stage", "source", "project", "owner", "temp", "due", "from", "to")}
    if f["q"]:
        digits = normalize_phone(f["q"])
        cond = Q(name__icontains=f["q"]) | Q(email__icontains=f["q"]) | Q(area__icontains=f["q"])
        if digits:
            cond |= Q(phone__contains=digits) | Q(alt_phone__contains=digits)
        queryset = queryset.filter(cond)
    if f["stage"] == "open":
        queryset = queryset.filter(stage__in=Lead.OPEN_STAGES)
    elif f["stage"]:
        queryset = queryset.filter(stage=f["stage"])
    if f["source"]:
        queryset = queryset.filter(source=f["source"])
    if f["project"]:
        queryset = queryset.filter(project_id=f["project"])
    if f["owner"] == "none":
        queryset = queryset.filter(owner__isnull=True)
    elif f["owner"]:
        queryset = queryset.filter(owner_id=f["owner"])
    if f["temp"]:
        queryset = queryset.filter(temperature=f["temp"])
    now = timezone.now()
    if f["due"] == "overdue":
        queryset = queryset.filter(stage__in=Lead.OPEN_STAGES, next_followup__lt=now)
    elif f["due"] == "today":
        queryset = queryset.filter(stage__in=Lead.OPEN_STAGES, next_followup__date=timezone.localdate())
    elif f["due"] == "untouched":
        queryset = queryset.filter(stage__in=Lead.OPEN_STAGES, first_response_at__isnull=True)
    for key, lookup in (("from", "created_at__date__gte"), ("to", "created_at__date__lte")):
        if f[key]:
            try:
                queryset = queryset.filter(**{lookup: datetime.strptime(f[key], "%Y-%m-%d").date()})
            except ValueError:
                f[key] = ""
    return queryset, f


def filter_options(user):
    return {
        "stages": Lead.STAGES,
        "sources": Lead.SOURCES,
        "temps": Lead.TEMPERATURES,
        "projects": Project.objects.all(),
        "owners": visible_users(user),
    }


@role_required(*SALES_ROLES)
def lead_list(request):
    leads, f = filter_leads(request, visible_leads(request.user).select_related("owner", "project"))
    sort = request.GET.get("sort", "new")
    leads = leads.order_by(
        {"followup": "next_followup", "updated": "-updated_at", "name": "name"}.get(sort, "-created_at")
    )
    page = Paginator(leads, 25).get_page(request.GET.get("page"))
    return render(
        request,
        "crm/leads/list.html",
        {"page": page, "f": f, "sort": sort, "total": page.paginator.count, "now": timezone.now(),
         "assignees": assignable_users(request.user), **filter_options(request.user)},
    )


@role_required(*SALES_ROLES)
def lead_board(request):
    leads, f = filter_leads(request, visible_leads(request.user).select_related("owner", "project"))
    since = timezone.now() - timedelta(days=BOARD_CLOSED_DAYS)
    leads = leads.exclude(stage__in=(Lead.BOOKED, Lead.LOST), stage_changed_at__lt=since)
    columns = {key: {"key": key, "label": label, "leads": []} for key, label in Lead.STAGES}
    temp_order = Case(When(temperature="hot", then=0), When(temperature="warm", then=1), default=2)
    for lead in leads.order_by(temp_order, "next_followup", "-created_at")[:600]:
        columns[lead.stage]["leads"].append(lead)
    return render(
        request,
        "crm/leads/board.html",
        {"columns": list(columns.values()), "f": f, "now": timezone.now(), "lost_reasons": Lead.LOST_REASONS,
         "closed_days": BOARD_CLOSED_DAYS, **filter_options(request.user)},
    )


@role_required(*SALES_ROLES)
def lead_create(request):
    form = LeadForm(request.POST or None, user=request.user)
    if request.method == "POST" and form.is_valid():
        data = dict(form.cleaned_data)
        try:
            lead, created = capture_lead(data, user=request.user, channel="Manual entry")
        except ValidationError as exc:
            form.add_error(None, "; ".join(exc.messages))
        else:
            if created:
                messages.success(request, f"Lead {lead.name} created and assigned to {lead.owner.get_full_name() if lead.owner else 'nobody'}.")
                return redirect("lead_detail", pk=lead.pk)
            if can_view_lead(request.user, lead):
                messages.warning(request, f"This number already exists as {lead.name}. The re-enquiry was recorded on the existing lead.")
                return redirect("lead_detail", pk=lead.pk)
            messages.warning(request, f"This number already exists and is handled by "
                                      f"{lead.owner.get_full_name() if lead.owner else 'another team'}. "
                                      "The re-enquiry was recorded and they have been notified.")
            return redirect("lead_list")
    return render(request, "crm/leads/form.html", {"form": form, "geo": geo_data()})


@role_required(*SALES_ROLES)
def lead_check_phone(request):
    phone = normalize_phone(request.GET.get("phone", ""))
    code = request.GET.get("code") or DEFAULT_DIAL
    phone = strip_dial_code(phone, code)
    if len(phone) < 6:
        return JsonResponse({"exists": False})
    lead = Lead.objects.select_related("owner").filter(phone_code=code, phone=phone).first()
    if not lead:
        return JsonResponse({"exists": False})
    data = {"exists": True, "owner": lead.owner.get_full_name() if lead.owner else "Unassigned"}
    if can_view_lead(request.user, lead):
        data.update(name=lead.name, stage=lead.get_stage_display(), url=reverse("lead_detail", args=[lead.pk]))
    return JsonResponse(data)


@role_required(*SALES_ROLES)
def lead_detail(request, pk):
    lead = get_visible_lead(request, pk)
    initial = {"temperature": lead.temperature, "kind": "call"}
    if lead.next_followup and lead.next_followup > timezone.now():
        initial["next_followup"] = timezone.localtime(lead.next_followup)
    return render(
        request,
        "crm/leads/detail.html",
        {
            "lead": lead,
            "activity_form": ActivityForm(initial=initial),
            "lost_form": LostForm(),
            "reassign_form": ReassignForm(user=request.user),
            "activities": lead.activities.select_related("created_by")[:100],
            "visits": lead.site_visits.select_related("project", "executive").order_by("-scheduled_at"),
            "assignments": lead.assignments.select_related("from_user", "to_user", "by_user")[:10],
            "previews": lead_previews(lead),
            "stage_list": [(k, v, Lead.STAGE_RANK[k]) for k, v in Lead.STAGES if k != Lead.LOST],
            "booking": getattr(lead, "booking", None) if lead.stage == Lead.BOOKED else None,
        },
    )


@role_required(*SALES_ROLES)
def lead_edit(request, pk):
    lead = get_visible_lead(request, pk)
    form = LeadForm(request.POST or None, instance=lead, user=request.user)
    if request.method == "POST" and form.is_valid():
        form.save()
        if form.changed_data:
            audit(request.user, "update", lead, f"{lead.name}: changed {', '.join(form.changed_data)}")
        messages.success(request, "Lead updated.")
        return redirect("lead_detail", pk=pk)
    return render(request, "crm/leads/form.html", {"form": form, "lead": lead, "geo": geo_data()})


@role_required(ADMIN)
@require_POST
def lead_delete(request, pk):
    lead = get_object_or_404(Lead, pk=pk)
    audit(request.user, "delete", lead, f"Deleted lead {lead.name} ({lead.full_phone})")
    lead.delete()
    messages.success(request, f"Lead {lead.name} deleted.")
    return redirect("lead_list")


@role_required(*SALES_ROLES)
@require_POST
def lead_log_activity(request, pk):
    lead = get_visible_lead(request, pk)
    form = ActivityForm(request.POST)
    if not form.is_valid():
        messages.error(request, "Could not save: " + "; ".join(e for errs in form.errors.values() for e in errs))
        return redirect("lead_detail", pk=pk)
    d = form.cleaned_data
    log_activity(lead, d["kind"], request.user, note=d["note"], outcome=d["outcome"])
    lead.refresh_from_db()
    lead.next_followup = d["next_followup"] if lead.is_open else None
    lead.temperature = d["temperature"]
    lead.save(update_fields=["next_followup", "temperature", "updated_at"])
    messages.success(request, "Activity logged." + (
        f" Next follow-up: {timezone.localtime(d['next_followup']):%d %b, %I:%M %p}." if d["next_followup"] else ""))
    return redirect("lead_detail", pk=pk)


@role_required(*SALES_ROLES)
@require_POST
def lead_stage(request, pk):
    lead = get_visible_lead(request, pk)
    stage = request.POST.get("stage", "")
    wants_json = request.headers.get("x-requested-with") == "fetch"

    def reply(ok, msg, url=None, status=200):
        if wants_json:
            return JsonResponse({"ok": ok, "message": msg, "redirect": url}, status=status)
        (messages.success if ok else messages.error)(request, msg)
        return redirect(url or reverse("lead_detail", args=[pk]))

    if stage == Lead.BOOKED and lead.stage != Lead.BOOKED:
        return reply(False, "Fill in the booking details to mark this lead as booked.",
                     reverse("booking_create", args=[pk]), status=409)
    if stage == Lead.SV_SCHEDULED and not lead.site_visits.filter(status="scheduled").exists():
        return reply(False, "Schedule the site visit with date and time.",
                     reverse("visit_schedule", args=[pk]), status=409)
    if lead.stage == Lead.BOOKED:
        return reply(False, "Booked leads can't be moved. Contact the admin to cancel a booking.", status=400)
    if stage == Lead.LOST:
        form = LostForm(request.POST)
        if not form.is_valid():
            return reply(False, "Please choose why the lead was lost.", status=400)
        set_stage(lead, stage, request.user, form.cleaned_data["reason"], form.cleaned_data["note"])
    else:
        try:
            set_stage(lead, stage, request.user)
        except ValueError as exc:
            return reply(False, str(exc), status=400)
    return reply(True, f"{lead.name} moved to {lead.get_stage_display()}.")


@role_required(ADMIN, MANAGER, LEADER)
@require_POST
def lead_reassign(request, pk):
    lead = get_visible_lead(request, pk)
    form = ReassignForm(request.POST, user=request.user)
    if form.is_valid():
        to = form.cleaned_data["to_user"]
        assign(lead, to, by_user=request.user, reason=form.cleaned_data["reason"] or "Manual reassignment")
        messages.success(request, f"Lead reassigned to {to.get_full_name()}.")
    else:
        messages.error(request, "Choose who to reassign the lead to.")
    return redirect("lead_detail", pk=pk)


@role_required(ADMIN, MANAGER, LEADER)
@require_POST
def lead_bulk_assign(request):
    to = assignable_users(request.user).filter(pk=request.POST.get("to_user")).first()
    ids = request.POST.getlist("ids")
    if not to or not ids:
        messages.error(request, "Select leads and a person to assign them to.")
        return redirect(request.POST.get("next") or "lead_list")
    leads = visible_leads(request.user).filter(pk__in=ids)
    for lead in leads:
        assign(lead, to, by_user=request.user, reason="Bulk reassignment")
    messages.success(request, f"{leads.count()} lead(s) assigned to {to.get_full_name()}.")
    nxt = request.POST.get("next", "")
    return redirect(nxt if nxt.startswith("/") else "lead_list")


# ---------------------------------------------------------------- import / export

IMPORT_COLUMNS = ["Name", "Country Code", "Phone", "Alternate Phone", "Email", "Country", "City", "Area",
                  "Budget", "Currency", "Property Type", "Project", "Source", "Notes"]


def _choice(value, choices, default=""):
    text = str(value or "").strip().lower()
    if not text:
        return default
    for key, label in choices:
        if text in (key.lower(), label.lower()):
            return key
    raise ValueError(f"'{value}' is not a valid option")


def _read_rows(upload):
    if upload.name.lower().endswith(".csv"):
        try:
            return list(csv.reader(io.StringIO(upload.read().decode("utf-8-sig"))))
        except UnicodeDecodeError as exc:
            raise ValueError("CSV must be UTF-8 encoded.") from exc
    try:
        wb = openpyxl.load_workbook(upload, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Could not read the Excel file.") from exc
    try:
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


@role_required(*MGMT)
def lead_import(request):
    form = ImportForm(request.POST or None, request.FILES or None)
    result = None
    if request.method == "POST" and form.is_valid():
        try:
            rows = _read_rows(form.cleaned_data["file"])
            if not rows:
                raise ValueError("The file is empty.")
            header = [_cell(h).lower() for h in rows[0]]
            index = {c: header.index(c.lower()) for c in IMPORT_COLUMNS if c.lower() in header}
            if "Name" not in index or "Phone" not in index:
                raise ValueError("The file needs at least 'Name' and 'Phone' columns.")
            if len(rows) > 5001:
                raise ValueError("Maximum 5,000 rows per import.")
        except ValueError as exc:
            form.add_error("file", str(exc))
        else:
            projects = {p.name.lower(): p for p in Project.objects.all()}
            result = {"created": 0, "duplicates": 0, "errors": []}
            for number, row in enumerate(rows[1:], start=2):
                def get(col, row=row):
                    i = index.get(col)
                    return _cell(row[i]) if i is not None and i < len(row) else ""

                if not any(_cell(c) for c in row):
                    continue
                try:
                    project_name = get("Project")
                    project = projects.get(project_name.lower()) if project_name else None
                    project_other = project_name[:100] if project_name and not project else ""
                    country = country_code(get("Country")) if get("Country") else DEFAULT_COUNTRY
                    if not country:
                        raise ValueError(f"Unknown country '{get('Country')}'")
                    dial = "+" + normalize_phone(get("Country Code")) if get("Country Code") else dial_for(country)
                    currency = (get("Currency") or BASE_CURRENCY).upper()
                    if currency not in CURRENCIES:
                        raise ValueError(f"Currency must be one of {', '.join(CURRENCIES)}")
                    budget = get("Budget").replace(",", "")
                    try:
                        source, source_other = _choice(get("Source"), Lead.SOURCES, "other"), ""
                    except ValueError:
                        source, source_other = "other", get("Source")[:60]
                    try:
                        ptype, ptype_other = _choice(get("Property Type"), Lead.PROPERTY_TYPES), ""
                    except ValueError:
                        ptype, ptype_other = "other", get("Property Type")[:60]
                    data = {
                        "name": get("Name"), "phone_code": dial, "phone": strip_dial_code(normalize_phone(get("Phone")), dial),
                        "alt_phone": normalize_phone(get("Alternate Phone")), "email": get("Email"),
                        "country": country, "city": get("City"), "area": get("Area"),
                        "budget_amount": to_base(budget, currency) if budget else None,
                        "property_type": ptype, "property_type_other": ptype_other,
                        "project": project, "project_other": project_other,
                        "source": source, "source_other": source_other, "notes": get("Notes"),
                    }
                    if budget and data["budget_amount"] is None:
                        raise ValueError(f"Budget '{budget}' is not a number")
                    if not data["name"] or len(data["phone"]) < 6:
                        raise ValueError("Name and a valid phone are required")
                    _, created = capture_lead(data, user=request.user, channel="Excel import")
                    result["created" if created else "duplicates"] += 1
                except (ValueError, ValidationError) as exc:
                    msg = "; ".join(exc.messages) if isinstance(exc, ValidationError) else str(exc)
                    result["errors"].append((number, get("Name") or "-", msg))
            audit(request.user, "import", Lead(), f"Imported {result['created']} leads, {result['duplicates']} re-enquiries")
            messages.success(request, f"Import finished: {result['created']} new, {result['duplicates']} re-enquiries, "
                                      f"{len(result['errors'])} errors.")
    return render(request, "crm/leads/import.html", {"form": form, "result": result, "columns": IMPORT_COLUMNS})


@role_required(*MGMT)
def lead_import_template(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(IMPORT_COLUMNS)
    ws.append(["Ahmed Al-Mansoori", "+974", "55123456", "", "ahmed@example.com", "Qatar", "Lusail", "Marina District",
               "1800000", "QAR", "2 Bedroom", Project.objects.values_list("name", flat=True).first() or "",
               "Property Finder", "Wants sea view"])
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="mvhomes-lead-import-template.xlsx"'
    return response


@role_required(*MGMT)
def lead_export(request):
    leads, _ = filter_leads(request, visible_leads(request.user).select_related("owner", "project"))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(["Name", "Phone", "Email", "Country", "City", "Area", "Project", "Property type", f"Budget ({BASE_CURRENCY})",
               "Source", "Stage", "Temperature", "Owner", "Next follow-up", "Lost reason", "Enquiries", "Created"])
    count = 0
    for lead in leads.order_by("-created_at").iterator():
        ws.append([
            lead.name, lead.full_phone, lead.email, lead.get_country_display(), lead.city, lead.area,
            lead.project_label, lead.get_property_type_display(),
            float(lead.budget_amount) if lead.budget_amount is not None else "", lead.get_source_display(),
            lead.get_stage_display(), lead.get_temperature_display(),
            lead.owner.get_full_name() if lead.owner else "",
            timezone.localtime(lead.next_followup).strftime("%d-%m-%Y %H:%M") if lead.next_followup else "",
            lead.get_lost_reason_display(), lead.enquiry_count, timezone.localtime(lead.created_at).strftime("%d-%m-%Y %H:%M"),
        ])
        count += 1
    audit(request.user, "export", Lead(), f"Exported {count} leads")
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="mvhomes-leads-{timezone.localdate():%Y%m%d}.xlsx"'
    return response
